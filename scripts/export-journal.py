#!/usr/bin/env python3
"""Експорт Google-таблиці журналу (xlsx) у data/journal.json для prisma/import-real.ts.
Використання: python3 scripts/export-journal.py "<шлях до xlsx>"  (результат у data/, не комітиться)."""
import json, re, sys
import openpyxl

src = sys.argv[1]
wb = openpyxl.load_workbook(src, data_only=True)
DATE = re.compile(r"^(\d{2})\.(\d{2})\.(\d{4})(?:\s+(\d))?$")
out = []
for ws in wb:
    cols = []
    for c in ws[1][2:]:
        raw = c.value
        if raw is None:
            continue
        h = " ".join(str(raw).split())
        m = DATE.match(h)
        if m:
            d, mo, y, dup = m.groups()
            cols.append({"col": c.column, "kind": "date", "date": f"{y}-{mo}-{d}", "dup": int(dup) if dup else 1})
        elif h.lower().startswith("практична"):
            cols.append({"col": c.column, "kind": "practical", "title": h})
        elif h.lower().startswith("модуль"):
            cols.append({"col": c.column, "kind": "module", "title": h})
        else:
            cols.append({"col": c.column, "kind": "skip", "title": h})
    students = []
    for row in ws.iter_rows(min_row=2):
        name = row[1].value
        if not name or str(name).startswith("Куратор"):
            continue
        has_num = isinstance(row[0].value, (int, float))
        has_data = any(c.value is not None for c in row[2:])
        if not (has_num or has_data):
            continue
        cells = {}
        for col in cols:
            v = row[col["col"] - 1].value
            if v is not None:
                cells[str(col["col"])] = v
        students.append({"name": str(name), "cells": cells})
    out.append({"sheet": ws.title, "columns": cols, "students": students})
json.dump(out, open("data/journal.json", "w"), ensure_ascii=False, indent=1)
print({s["sheet"]: len(s["students"]) for s in out})
